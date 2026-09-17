import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_documentation_docx(output_path):
    doc = docx.Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Color Palette
    PRIMARY_COLOR = RGBColor(79, 70, 229)    # Indigo 600
    SECONDARY_COLOR = RGBColor(30, 41, 59)  # Slate 800
    TEXT_COLOR = RGBColor(51, 65, 85)        # Slate 700
    MUTED_COLOR = RGBColor(100, 116, 139)   # Slate 500

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("HackBridge — System Documentation & Manual")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Enterprise Multi-Tenant Hackathon Lifecycle & Direct Recruitment Platform\n"
                           "Pilot Tenant: Maharaja Institute of Technology Thandavapura (MITT) | Version 2.0")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = MUTED_COLOR

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(12)
    div_run = div_p.add_run("―" * 58)
    div_run.font.color.rgb = RGBColor(226, 232, 240)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        return p

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            bp = p.add_run(bold_prefix)
            bp.font.name = "Calibri"
            bp.font.size = Pt(10.5)
            bp.font.bold = True
            bp.font.color.rgb = SECONDARY_COLOR
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = TEXT_COLOR
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            bp = p.add_run(bold_prefix)
            bp.font.name = "Calibri"
            bp.font.size = Pt(10)
            bp.font.bold = True
            bp.font.color.rgb = SECONDARY_COLOR
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_COLOR
        return p

    def add_callout(text, title="NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        t_run = p.add_run(f"[{title}] ")
        t_run.font.name = "Calibri"
        t_run.font.size = Pt(9.5)
        t_run.font.bold = True
        t_run.font.color.rgb = PRIMARY_COLOR
        c_run = p.add_run(text)
        c_run.font.name = "Calibri"
        c_run.font.size = Pt(9.5)
        c_run.font.color.rgb = SECONDARY_COLOR
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Executive Summary
    add_heading_1("1. Executive Summary")
    add_body(
        "HackBridge is a purpose-built, multi-stakeholder hackathon lifecycle and recruitment SaaS platform designed "
        "to solve operational fragmentation in university-level innovation challenges. Pilot deployed for Maharaja Institute of "
        "Technology Thandavapura (MITT), HackBridge bridges collegiate talent, academic governance, and corporate tech recruitment "
        "within a single secure, white-labeled system."
    )
    add_bullet(" Institutional oversight of hackathon stages, custom rubrics, AI triage, and audit trails.", "College Administrators:")
    add_bullet(" Team formation, invite codes, problem statement intake, deliverables submission, and verified talent profiles.", "Student Innovators:")
    add_bullet(" Double-blind scoring interface with real-time multi-criteria rubric sliders, weighted scoring, and COI recusal.", "Evaluators & Judges:")
    add_bullet(" Industrial challenge authoring, team tracking, talent pool scouting, and 1-click hiring outreach.", "Company Recruiters:")

    # 2. System Architecture & Tech Stack
    add_heading_1("2. System Architecture & Technical Specifications")
    add_body(
        "HackBridge leverages a decoupled cloud architecture combining a high-performance React 18 single-page application (SPA) "
        "with a Supabase PostgreSQL backend, fortified by server-side Row-Level Security (RLS) policies and GoTrue authentication."
    )

    # Tech Stack Table
    t_tech = doc.add_table(rows=1, cols=4)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer", "Technology", "Version", "Key Functional Purpose"]
    hdr_cells = t_tech.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "4F46E5")
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

    tech_rows = [
        ("Frontend", "React + Vite", "18.3 / 5.4", "Component-driven SPA with sub-second HMR and tree-shaken production bundle"),
        ("Language", "TypeScript", "5.5.3", "Strict static typing, null safety, and database schema interface parity"),
        ("Styling", "Tailwind CSS", "3.4.1", "Responsive design system with institutional CSS custom property theming"),
        ("Backend Database", "PostgreSQL 15", "Supabase", "Multi-tenant relational database with 14 applied idempotent migrations"),
        ("Access Control", "PostgreSQL RLS", "Native", "Declarative row-level tenant and role isolation policies"),
        ("Auth Engine", "Supabase GoTrue", "Native", "JWT token lifecycle, role claims, and fail-closed session guards"),
        ("Object Storage", "Supabase Storage", "Native S3", "5 buckets for banners, specs, deliverables, decks, and student CVs"),
        ("Deployment", "Docker + Nginx", "Multi-stage", "Alpine-based containerization with SPA HTML5 fallback proxy routing"),
    ]

    for row_data in tech_rows:
        row = t_tech.add_row()
        for idx, val in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = val
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            p.runs[0].font.color.rgb = TEXT_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. Security & Multi-Tenancy
    add_heading_1("3. Multi-Tenancy & Security Hardening")
    add_body(
        "Tenant isolation is strictly maintained at both the application gateway and database layer. "
        "Every tenant possesses a record in public.tenants with domain bindings and brand assets."
    )
    add_bullet(" Client input can never supply or modify tenant_id; resolution is derived server-side via email domain or admin binding.", "Server-Side Tenant Resolution:")
    add_bullet(" Enforces 7 hierarchical roles: super_admin, college_admin, committee_member, evaluator, company_rep, mentor, student.", "Hierarchical RBAC:")
    add_bullet(" Prevents any user from modifying their own role, tenant_id, or active status.", "Field Protection Trigger:")
    add_bullet(" In production, any missing Supabase configuration immediately halts workspace rendering.", "Fail-Closed Protection:")

    # 4. 12-Phase Full Specification Breakdown
    add_heading_1("4. 12-Phase Complete Specification Parity")
    add_body("HackBridge delivers 100% specification parity across all 12 institutional phases:")

    phases = [
        ("Phase 1 & 1.5: Multi-Tenant Core & Security", "Schema tenant isolation, GoTrue authentication, role whitelisting, and field protection triggers."),
        ("Phase 2A & 2B: Hackathon Lifecycle Engine", "7 database-enforced lifecycle states (draft to archived), prize pool tiers, and configurable rubrics."),
        ("Phase 2C: Corporate Partner Directory", "Company profile registry, industry domain classification, sponsor tier tracking, and verification desk."),
        ("Phase 3A & 3B: Problem Statement Intake & Review", "Industrial challenge authoring, review workflow, dataset attachments, and student challenge publication."),
        ("Phase 4: Team Formation & Capacity Enforcer", "Team creation, unique invite codes (e.g. MITT-9X4K), capacity limits (2-4), and challenge locking."),
        ("Phase 5: Multi-Format Submission Engine", "Project deliverables intake (GitHub repo, live demo, video, presentation deck, architecture abstract)."),
        ("Phase 6: Automated AI Pre-Screening Pipeline", "Heuristic triage engine computing relevance, completeness, and innovation scores (9.6/10) with flags."),
        ("Phase 7: Double-Blind Rubric Scoring Workspace", "Masked team dossiers, 4-criterion dynamic sliders, real-time weighted scores, and COI recusal."),
        ("Phase 8: Championship Dynamic Leaderboard", "Live podium rankings (Gold, Silver, Bronze), evaluator consensus metrics, and public ceremony view."),
        ("Phase 9: Verified Student Talent Portfolios", "Verified credentials, CGPA, hackathon award badges, interactive skills radar, and public portfolio URLs."),
        ("Phase 10: Recruiter Scouting & Direct Hiring", "Searchable student talent pool with filters (skills, CGPA) and 1-click Express Hiring Interest offers."),
        ("Phase 11: Compliance Audit Logs & Notifications", "Tamper-evident institutional audit trail (actor, action, IP) and real-time in-app notification center."),
        ("Phase 12: Production Infrastructure & Storage", "Supabase Storage buckets, Docker multi-stage build, Nginx Alpine proxy, and Redis/BullMQ queue design."),
    ]

    for p_title, p_desc in phases:
        add_bullet(f" {p_desc}", p_title)

    # 5. Stakeholder Manuals
    add_heading_1("5. Comprehensive Stakeholder User Manual")

    add_heading_2("5.1 College Administrator Guide (/admin)")
    add_bullet(" Create new hackathons, configure dates, rules, team bounds, and multi-criteria rubrics.", "Hackathon Orchestration:")
    add_bullet(" Verify company credentials and approve sponsorship participation.", "Corporate Partner Approval:")
    add_bullet(" Review, approve, or request revisions on industrial problem statements.", "Problem Statement Review:")
    add_bullet(" Inspect automated submission scores, repository health, and trigger batch triage.", "AI Pre-Screening Console:")
    add_bullet(" Review score distributions, resolve variance anomalies, and confirm prize awards.", "Judging Deliberation & Results:")
    add_bullet(" Access immutable audit logs for NAAC/NBA academic accreditation compliance.", "Compliance Audit Trail:")

    add_heading_2("5.2 Student Innovator Guide (/student)")
    add_bullet(" View registered team ('NeuralByte Innovations'), copy invite code ('MITT-9X4K'), and manage member roles.", "Team Workspace:")
    add_bullet(" Browse published challenges sponsored by Bosch, Zerodha, Philips, Infosys, and AgriTech.", "Problem Selection:")
    add_bullet(" Submit repo, live demo, video walkthrough, pitch deck, and inspect AI 9.6/10 readiness score.", "Deliverables Submission:")
    add_bullet(" Showcase verified 9.4 CGPA, Grand Champion badge, skills radar, and download verified resume.", "Talent Portfolio:")
    add_bullet(" Review corporate job and internship offers (e.g. Bosch ₹12-15 LPA CTC, Zerodha ₹60k/mo).", "Career Opportunities:")

    add_heading_2("5.3 Evaluator / Industry Judge Guide (/evaluator)")
    add_bullet(" View assigned submissions with status badges (Completed, Pending, Recused) and filter tabs.", "Assigned Submissions Queue:")
    add_bullet(" Double-blind masked dossier (left) with direct links to repo, demo, and slides.", "Scoring Workspace:")
    add_bullet(" Adjust 4-criterion sliders (Innovation, Complexity, Feasibility, Presentation) with live weighted score.", "Dynamic Rubric Sliders:")
    add_bullet(" Seamlessly toggle between assigned projects right inside the scoring workspace.", "Review Entry Switcher:")
    add_bullet(" One-click recusal protocol to declare conflict of interest and re-route submissions.", "COI Declaration:")

    add_heading_2("5.4 Corporate Recruiter Guide (/company)")
    add_bullet(" Submit technical challenges with target domains, bounty prizes, and dataset specs.", "Challenge Authoring:")
    add_bullet(" Search high-performing students by technical skill tags, CGPA, and hackathon distinctions.", "Talent Pool Scouting:")
    add_bullet(" Dispatch job/internship offers with compensation packages directly to students.", "1-Click Hiring Outreach:")

    # 6. Seeded Demo Presentation Dataset
    add_heading_1("6. Demonstration Dataset (MITT NIH 2026)")
    add_body(
        "The platform includes a cohesive, realistic dataset centered on the MITT National Innovation Hackathon 2026 "
        "allowing comprehensive live panel demonstrations without empty states."
    )

    # Seeded Teams Table
    t_teams = doc.add_table(rows=1, cols=5)
    t_teams.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Team Name", "Invite Code", "Problem Statement", "AI Score", "Championship Rank"]
    for i, h in enumerate(t_headers):
        c = t_teams.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "4F46E5")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

    team_data = [
        ("NeuralByte Innovations", "MITT-9X4K", "Bosch Autonomous Traffic Sync", "9.6 / 10", "1st Place (Gold — 95.5)"),
        ("QuantEdge AI", "MITT-7B2Y", "Zerodha Graph Mempool Interceptor", "9.4 / 10", "2nd Place (Silver — 93.0)"),
        ("AgriSense IoT", "MITT-5K9M", "Karnataka AgriTech Soil LoRa Mesh", "9.2 / 10", "3rd Place (Bronze — 91.5)"),
        ("HealthPulse Edge", "MITT-3W8L", "Philips Rural PHC AI Diagnostics", "8.8 / 10", "4th Place (88.0)"),
        ("CyberShield Zero", "MITT-2P4N", "Infosys eBPF Ransomware Interceptor", "8.7 / 10", "5th Place (86.5)"),
    ]

    for row_data in team_data:
        row = t_teams.add_row()
        for idx, val in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = val
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            p.runs[0].font.color.rgb = TEXT_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 7. Deployment Guide
    add_heading_1("7. Deployment, Docker & Operations Guide")
    add_body("HackBridge supports local development execution as well as containerized cloud deployment:")
    add_bullet(" npm install && npm run dev (accessible at http://localhost:5173/)", "Local Development:")
    add_bullet(" npm run build (verifies strict TypeScript compilation and builds minified bundle in dist/)", "Production Compilation:")
    add_bullet(" docker-compose up --build -d (launches Alpine Nginx container serving the optimized SPA on port 80)", "Docker Deployment:")

    # Save document
    doc.save(output_path)
    print(f"Documentation DOCX successfully generated at: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "HackBridge_Complete_Documentation.docx")
    create_documentation_docx(out_file)
